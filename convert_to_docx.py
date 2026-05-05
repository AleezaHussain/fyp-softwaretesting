"""
Converts Test_Strategy_Document.md to a formatted .docx file.
Font: Times New Roman, Size 12, Line Spacing 1.15, Justified
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm
import re

# ── helpers ──────────────────────────────────────────────────────────────────

def set_font(run, bold=False, size=12, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_para_format(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
    fmt = para.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = Pt(13.8)   # 1.15 × 12pt

def add_heading(doc, text, level=1):
    para = doc.add_paragraph()
    set_para_format(para, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4)
    run = para.add_run(text)
    sizes = {1: 16, 2: 14, 3: 12}
    set_font(run, bold=True, size=sizes.get(level, 12))
    return para

def add_body(doc, text, bold=False):
    para = doc.add_paragraph()
    set_para_format(para)
    run = para.add_run(text)
    set_font(run, bold=bold)
    return para

def add_bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    set_para_format(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    # strip leading "- " or "• "
    text = re.sub(r"^[-•]\s*", "", text)
    # handle inline bold **...**
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        run = para.add_run(part)
        set_font(run, bold=(i % 2 == 1))
    return para

def style_table(table):
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = Pt(13.8)
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

def shade_row(row, hex_color="D9E1F2"):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

def bold_row(row):
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True

def parse_inline(para, text):
    """Add runs to para handling **bold** and plain text."""
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        run = para.add_run(part)
        set_font(run, bold=(i % 2 == 1))

# ── main builder ─────────────────────────────────────────────────────────────

def build_docx():
    with open("Test_Strategy_Document.md", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin   = Cm(3.17)
        section.right_margin  = Cm(3.17)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # ── H1 (#)
        if line.startswith("# ") and not line.startswith("## "):
            add_heading(doc, line[2:], level=1)
            i += 1

        # ── H2 (##)
        elif line.startswith("## ") and not line.startswith("### "):
            add_heading(doc, line[3:], level=2)
            i += 1

        # ── H3 (###)
        elif line.startswith("### "):
            add_heading(doc, line[4:], level=3)
            i += 1

        # ── H4 (####)
        elif line.startswith("#### "):
            add_heading(doc, line[5:], level=3)
            i += 1

        # ── Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1

        # ── Table (starts with |)
        elif line.startswith("|"):
            # collect all table lines
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i].rstrip("\n"))
                i += 1

            # filter out separator rows (---|---)
            data_rows = [r for r in table_lines if not re.match(r"^\|[-| :]+\|$", r)]

            if not data_rows:
                continue

            # parse cells
            def parse_row(r):
                cells = [c.strip() for c in r.strip("|").split("|")]
                return cells

            rows = [parse_row(r) for r in data_rows]
            cols = max(len(r) for r in rows)

            table = doc.add_table(rows=len(rows), cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for ri, row_data in enumerate(rows):
                row = table.rows[ri]
                for ci, cell_text in enumerate(row_data):
                    if ci >= cols:
                        break
                    cell = row.cells[ci]
                    cell.text = ""
                    para = cell.paragraphs[0]
                    parse_inline(para, cell_text)
                    for run in para.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)
                        if ri == 0:
                            run.font.bold = True
                    para.paragraph_format.space_before = Pt(2)
                    para.paragraph_format.space_after  = Pt(2)
                    para.paragraph_format.line_spacing = Pt(13.8)

                if ri == 0:
                    shade_row(row, "2E74B5")
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif ri % 2 == 0:
                    shade_row(row, "D6E4F0")

            style_table(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # ── Bullet point
        elif line.startswith("- ") or line.startswith("* "):
            add_bullet(doc, line)
            i += 1

        # ── Numbered list
        elif re.match(r"^\d+\.\s", line):
            para = doc.add_paragraph(style="List Number")
            set_para_format(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            text = re.sub(r"^\d+\.\s*", "", line)
            parse_inline(para, text)
            for run in para.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
            i += 1

        # ── Bold line (**text**)
        elif line.startswith("**") and line.endswith("**") and len(line) > 4:
            add_body(doc, line.strip("*"), bold=True)
            i += 1

        # ── Empty line
        elif line.strip() == "":
            i += 1

        # ── Normal paragraph
        else:
            para = doc.add_paragraph()
            set_para_format(para)
            parse_inline(para, line)
            for run in para.runs:
                if not run.font.size:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
            i += 1

    doc.save("Test_Strategy_Document.docx")
    print("✅  Saved: Test_Strategy_Document.docx")

if __name__ == "__main__":
    build_docx()
