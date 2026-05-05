"""
Converts Test_Strategy_Document.md to Test_Strategy_Document.docx
Run: py convert_to_word.py
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

md_text = Path("Test_Strategy_Document.md").read_text(encoding="utf-8")
doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def add_table_from_md(doc, lines):
    rows = [l for l in lines if l.strip().startswith('|')]
    if len(rows) < 2:
        return
    # Remove separator row
    data_rows = [r for r in rows if not re.match(r'^\|[-| :]+\|$', r.strip())]
    if not data_rows:
        return
    parsed = []
    for row in data_rows:
        cells = [c.strip() for c in row.strip().strip('|').split('|')]
        parsed.append(cells)
    
    if not parsed:
        return
    
    cols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = 'Table Grid'
    
    for i, row in enumerate(parsed):
        for j, cell_text in enumerate(row):
            if j < cols:
                cell = table.cell(i, j)
                cell.text = cell_text
                # Header row bold
                if i == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                        # Header background
                        tc = cell._tc
                        tcPr = tc.get_or_add_tcPr()
                        shd = OxmlElement('w:shd')
                        shd.set(qn('w:val'), 'clear')
                        shd.set(qn('w:color'), 'auto')
                        shd.set(qn('w:fill'), '1F3864')
                        tcPr.append(shd)
                        for run in para.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()

lines = md_text.split('\n')
i = 0
table_buffer = []
in_table = False

while i < len(lines):
    line = lines[i]
    
    # Detect table start
    if line.strip().startswith('|'):
        in_table = True
        table_buffer.append(line)
        i += 1
        continue
    
    # Table ended
    if in_table and not line.strip().startswith('|'):
        add_table_from_md(doc, table_buffer)
        table_buffer = []
        in_table = False
    
    # Headings
    if line.startswith('# '):
        p = doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        p = doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        p = doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('#### '):
        p = doc.add_heading(line[5:].strip(), level=4)
    # Horizontal rule
    elif line.strip() == '---':
        doc.add_paragraph('─' * 60)
    # Bullet points
    elif line.strip().startswith('- '):
        doc.add_paragraph(line.strip()[2:], style='List Bullet')
    elif re.match(r'^\d+\. ', line.strip()):
        doc.add_paragraph(re.sub(r'^\d+\. ', '', line.strip()), style='List Number')
    # Bold inline
    elif line.strip():
        para = doc.add_paragraph()
        # Handle **bold** inline
        parts = re.split(r'\*\*(.*?)\*\*', line)
        for idx, part in enumerate(parts):
            run = para.add_run(part)
            if idx % 2 == 1:
                run.bold = True
    else:
        doc.add_paragraph()
    
    i += 1

# Handle trailing table
if table_buffer:
    add_table_from_md(doc, table_buffer)

doc.save("Test_Strategy_Document.docx")
print("✅ Saved: Test_Strategy_Document.docx")
